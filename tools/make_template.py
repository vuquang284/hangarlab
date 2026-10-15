"""Builds 'HangarLab template.xlsx' and template.js (base64 copy the app offers for download).
Run: python make_template.py   (needs openpyxl)"""
import base64, colorsys, os
from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

HERE = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.join(HERE, 'HangarLab template.xlsx')

# same order and hues as the app (app.js SUB_HUE)
SUBS = [('Structures', 212), ('3D printed', 84), ('Propulsion', 24), ('Power', 352), ('Avionics', 268), ('Sensors', 172),
        ('Comms', 196), ('Compute', 316), ('Payload', 140), ('Wiring', 50), ('Hardware', 32), ('Other', 220)]
RAILS = ['VBAT', '3.3V', '5V', '9V', '12V', '24V', 'Own supply', 'None']
BANDS = ['None', '433 MHz', '868 MHz', '915 MHz', '1.2/1.3 GHz', '2.4 GHz', '5.8 GHz']
CHEMS = ['LiPo', 'LiHV', 'Li-ion', 'LiFePO4']

ORANGE, INK, MUTED = 'D9530A', '1D1D1B', '6F6D65'
HEAD = PatternFill('solid', fgColor='2B2B28')
HFONT = Font(bold=True, color='FFFFFF')
EX_FONT = Font(italic=True, color=MUTED)
THIN = Border(bottom=Side(style='thin', color='DCDAD2'))


def hsl_hex(h, s, l):
    r, g, b = colorsys.hls_to_rgb(h / 360, l, s)
    return '%02X%02X%02X' % (round(r * 255), round(g * 255), round(b * 255))


def table(ws, headers, widths, rows, notes=None):
    ws.append(headers)
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
        c = ws.cell(row=1, column=i)
        c.fill, c.font, c.alignment = HEAD, HFONT, Alignment(vertical='center')
    ws.row_dimensions[1].height = 22
    for r in rows:
        ws.append(r)
        for c in ws[ws.max_row]:
            c.font, c.border = EX_FONT, THIN
    ws.freeze_panes = 'A2'
    if notes:
        for i, n in enumerate(notes, 1):
            if n:
                ws.cell(row=1, column=i).comment = None
    return ws


def dropdown(ws, col, source, rows=500):
    dv = DataValidation(type='list', formula1=source, allow_blank=True, showErrorMessage=False)
    ws.add_data_validation(dv)
    dv.add(f'{col}2:{col}{rows}')


wb = Workbook()

# ---------- Read me ----------
rd = wb.active
rd.title = 'Read me'
rd.column_dimensions['A'].width = 110
lines = [
    ('HangarLab component template', Font(bold=True, size=16, color=ORANGE)),
    ('Fill in the sheets, save, then in the app open Library → Import workbook. Rows are matched by Name: new names are added, existing ones updated.', None),
    ('', None),
    ('Sheets', Font(bold=True, size=12)),
    ('Components — avionics, sensors, radios, payloads, wiring, hardware… Pick a Subsystem from the dropdown; the cell takes the app’s color.', None),
    ('Motors — kV, winding resistance Rm (mΩ) and no-load current I0 turn on the motor model. Max current/power from the datasheet.', None),
    ('Props — diameter and pitch in inches. Performance data (APC .dat) is imported separately in the app.', None),
    ('Batteries — one row per pack: cells in series (S), in parallel (P), capacity per cell, continuous C and peak (burst) C. Pack count is set per design in the app.', None),
    ('ESCs — continuous and burst current, the cell range it accepts, and weight. Designs pick an ESC for the lift and cruise motors.', None),
    ('Airframes — type, then wingspan, wing area and length in meters for wings, or wheelbase (mm), layout and max prop size for multirotor frames; weight in grams.', None),
    ('Motor thrust data — thrust-stand tests, one row per measurement: Motor, Prop, Throttle, Thrust, Voltage, Current. Rows with the same Motor + Prop form one test; pick names from the Motors and Props sheets. (Older one-sheet-per-test "Thrust - …" sheets still import.)', None),
    ('', None),
    ('Rules', Font(bold=True, size=12)),
    ('Units are metric: grams, meters, m², volts, amps, watts, mAh. (The app can still display imperial.)', None),
    ('Keep the header row. Columns can be reordered; extra columns are ignored.', None),
    ('Gray italic rows are examples: overwrite or delete them.', None),
    ('Power: Avg W is typical draw, Peak W the worst case. Rail = the supply that powers the part (VBAT = straight from the flight pack).', None),
    ('Link = datasheet or product page. Buy link = where to order it. Unit cost is per item in your currency; the BOM report totals it.', None),
    ('Qty in Components is only a reference for your own BOM; the library stores one of each. Add parts to a design in the app.', None),
]
for text, font in lines:
    rd.append([text])
    if font:
        rd.cell(row=rd.max_row, column=1).font = font
    rd.cell(row=rd.max_row, column=1).alignment = Alignment(wrap_text=True, vertical='top')

# ---------- lists (hidden) ----------
ls = wb.create_sheet('Lists')
for i, (name, _) in enumerate(SUBS, 1):
    ls.cell(row=i, column=1, value=name)
for i, v in enumerate(RAILS, 1):
    ls.cell(row=i, column=2, value=v)
for i, v in enumerate(BANDS, 1):
    ls.cell(row=i, column=3, value=v)
for i, v in enumerate(CHEMS, 1):
    ls.cell(row=i, column=4, value=v)
ls.sheet_state = 'hidden'

# ---------- Components ----------
cp = wb.create_sheet('Components')
table(cp, ['Name', 'Subsystem', 'Weight (g)', 'Qty', 'Rail', 'Avg power (W)', 'Peak power (W)', 'RF Freq', 'Link', 'Notes', 'Unit cost', 'Buy link'],
      [36, 16, 12, 7, 12, 14, 15, 18, 42, 30, 11, 42], [
          ['Cube Orange+ Mini', 'Avionics', 53, 1, '5V', 2.5, 3, 'None', 'https://docs.cubepilot.org/', '', 395, 'https://irlock.com/products/the-cube-orange-mini-set'],
          ['Here4 GPS', 'Sensors', 60, 1, '5V', 1, 1.5, 'None', '', ''],
          ['RFD900x telemetry', 'Comms', 18, 1, '5V', 5, 5, '915 MHz', '', ''],
          ['ELRS RP4TD receiver', 'Comms', 3.4, 1, '5V', 0.5, 0.6, '2.4 GHz', '', ''],
          ['Jetson Orin Nano + Seeed A603', 'Compute', 128, 1, '12V', 15, 25, 'None', '', ''],
          ['SIYI A8 mini gimbal camera', 'Payload', 95, 1, '12V', 5, 8, 'None', '', ''],
          ['TBS Lucid 12S 90A ESC', 'Propulsion', 20, 4, 'None', 0, 0, 'None', '', ''],
          ['PDB & power modules', 'Power', 70, 1, 'None', 0, 0, 'None', '', ''],
          ['Silicone wire, connectors', 'Wiring', 250, 1, 'None', 0, 0, 'None', '', 'estimate'],
          ['M3 hardware kit', 'Hardware', 25, 1, 'None', 0, 0, 'None', '', ''],
          ['Carbon frame arms', 'Structures', 84, 4, 'None', 0, 0, 'None', '', ''],
          ['GPS mast (PETG)', '3D printed', 6.5, 1, 'None', 0, 0, 'None', '', ''],
      ])
dropdown(cp, 'B', '=Lists!$A$1:$A$%d' % len(SUBS))
dropdown(cp, 'E', '=Lists!$B$1:$B$%d' % len(RAILS))
dropdown(cp, 'H', '=Lists!$C$1:$C$%d' % len(BANDS))
for name, hue in SUBS:  # subsystem cell takes the app color
    sat = 0.12 if name == 'Other' else 0.3 if name == 'Hardware' else 0.62
    fill = PatternFill('solid', fgColor=hsl_hex(hue, sat, 0.86), bgColor=hsl_hex(hue, sat, 0.86))
    cp.conditional_formatting.add('B2:B500', FormulaRule(formula=[f'$B2="{name}"'], fill=fill, font=Font(bold=True, color=hsl_hex(hue, min(sat + 0.1, 1), 0.3))))

# ---------- Motors ----------
mo = wb.create_sheet('Motors')
table(mo, ['Name', 'kV', 'Rm (mΩ)', 'I0 (A)', 'I0 voltage (V)', 'Max current (A)', 'Max power (W)', 'Weight (g)', 'Link', 'Unit cost', 'Buy link'],
      [32, 8, 10, 9, 14, 16, 15, 12, 42, 11, 42], [
          ['T-Motor MN5008 340KV', 340, 89, 0.8, 10, 30, 700, 135, 'https://store.tmotor.com/', 110, 'https://store.tmotor.com/'],
          ['Hobbywing 4120 250KV', 250, '', '', '', 55.3, '', 309, '', '', ''],
      ])

# ---------- Motor thrust data (one row per measurement; motor/prop names match the Motors and Props sheets) ----------
mt = wb.create_sheet('Motor thrust data')
table(mt, ['Motor', 'Prop', 'Throttle (%)', 'Thrust (g)', 'Voltage (V)', 'Current (A)', 'Power (W)'], [30, 18, 13, 12, 12, 12, 12],
      [['Hobbywing 4120 250KV', 'APC 16x8E', t, g, 44.4, a] for t, g, a in [(33, 1029, 2.3), (42, 1642, 4.2), (51, 2272, 7.0), (60, 3108, 11.3), (72, 4488, 19.5), (81, 5560, 28.4), (90, 6575, 39.9), (100, 7694, 55.3)]]
      + [['T-Motor MN5008 340KV', 'APC 15x5.5MR', t, g, 22.2, a] for t, g, a in [(50, 1150, 4.1), (75, 2210, 10.6), (100, 3150, 20.3)]])
for r in range(2, mt.max_row + 1):
    mt.cell(row=r, column=7, value=f'=E{r}*F{r}').font = EX_FONT
dropdown(mt, 'A', '=Motors!$A$2:$A$500')
dropdown(mt, 'B', '=Props!$A$2:$A$500')

# ---------- Props ----------
pr = wb.create_sheet('Props')
table(pr, ['Name', 'Diameter (in)', 'Pitch (in)', 'Weight (g)', 'Link', 'Unit cost', 'Buy link'], [26, 14, 12, 12, 42, 11, 42], [
    ['APC 16x8E', 16, 8, 51.87, 'https://www.apcprop.com/', 12, 'https://www.apcprop.com/'],
    ['HQ 8x4.5x3', 8, 4.5, '', '', '', ''],
    ['APC 15x5.5MR', 15, 5.5, 42, '', '', ''],
])

# ---------- ESCs ----------
es = wb.create_sheet('ESCs')
table(es, ['Name', 'Brand', 'For (Multirotor / Plane / Either)', 'Continuous current (A)', 'Burst current (A)', 'Min cells (S)', 'Max cells (S)', 'Weight (g)', 'Link', 'Unit cost', 'Buy link'],
      [28, 14, 16, 20, 17, 13, 13, 11, 42, 11, 42], [
          ['Generic 6S 60A ESC', '', 'Either', 60, 80, 3, 6, 60, '', '', ''],
          ['Generic 12S 80A ESC', '', 'Multirotor', 80, 100, 6, 12, 90, '', '', ''],
      ])

# ---------- Batteries ----------
ba = wb.create_sheet('Batteries')
table(ba, ['Name', 'Chemistry', 'Cells S', 'Cells P', 'Cell capacity (mAh)', 'Pack weight (g)', 'Continuous C', 'Peak C', 'Link', 'Unit cost', 'Buy link'],
      [40, 12, 9, 9, 19, 16, 14, 10, 42, 11, 42], [
          ['Tattu 6S 7000 25C', 'LiPo', 6, 1, 7000, 870, 25, 50, '', '', ''],
          ['Upgrade Energy GREEN V2 6S1P 4Ah Amprius SA10', 'Li-ion', 6, 1, 4000, 305, 1.5, 4, 'https://www.upgradeenergytech.com/', '', 'https://www.upgradeenergytech.com/'],
      ])
dropdown(ba, 'B', '=Lists!$D$1:$D$%d' % len(CHEMS))

# ---------- Airframes ----------
af = wb.create_sheet('Airframes')
table(af, ['Name', 'Type (Fixed-wing / Multirotor / Hybrid VTOL)', 'Wingspan (m)', 'Wing area (m²)', 'Length (m)', 'Wheelbase (mm)', 'Layout (Quad X, Hexa, ...)', 'Max prop (in)', 'Weight (g)', 'Notes', 'Link', 'Unit cost', 'Buy link'], [30, 18, 14, 15, 12, 14, 16, 13, 12, 30, 42, 11, 42], [
    ['AtomRC Dolphin Pro', 'Fixed-wing', 0.845, 0.1696, 0.826, '', '', '', 790, '', '', '', ''],
    ['HEEWING T2 Cruza', 'Fixed-wing', 1.2, '', '', '', '', '', 1200, 'wing area not listed', '', '', ''],
    ['Example 500 mm quad frame', 'Multirotor', '', '', '', 500, 'Quad X', 13, 450, 'multirotor: wheelbase, layout and max prop instead of a wing', '', '', ''],
])

wb.move_sheet('Lists', offset=10)
wb.save(OUT)
with open(OUT, 'rb') as f:
    b64 = base64.b64encode(f.read()).decode()
with open(os.path.join(HERE, '..', 'hangarlab', 'template.js'), 'w', encoding='utf8') as f:
    f.write('// Generated by make_template.py: the component template workbook, offered for download in Library.\n')
    f.write(f'(typeof window !== "undefined" ? window : globalThis).UAS_TEMPLATE_B64 = "{b64}";\n')
print('wrote', OUT, len(b64) // 1024, 'KB base64')
